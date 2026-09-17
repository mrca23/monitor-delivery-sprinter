import sys, glob, os
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import warnings; warnings.filterwarnings('ignore')
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = sys.argv[1] if len(sys.argv) > 1 else max(glob.glob(os.path.join(BASE, 'input', '*.xlsx')), key=os.path.getmtime)
d = pd.read_excel(src)

# Definisi
d['F_TTD'] = (d['TTD Retur'] == 0) & d['Waktu TTD'].notna()          # TTD normal oleh penerima
d['F_BERMASALAH'] = d['Waktu Paket Bermasalah'].notna()
d['F_INVENTORI'] = d['Waktu Scan Inventory'].notna()
d['F_GAGAL'] = ~d['F_TTD']                                           # tidak TTD (retur / belum TTD)
d['F_GAGAL_NO_INV'] = d['F_GAGAL'] & ~d['F_INVENTORI']               # gagal kirim tapi tidak scan inventori
d['F_GAGAL_NO_BM'] = d['F_GAGAL'] & ~d['F_BERMASALAH']               # gagal kirim tanpa scan bermasalah

tgl = pd.to_datetime(d['Waktu Delivery']).dt.date
# Wajib scan = AWB tidak TTD (Total Delivery - Total TTD): wajib scan bermasalah & wajib scan inventori
d['F_GAGAL_BM'] = d['F_GAGAL'] & d['F_BERMASALAH']
d['F_GAGAL_INV'] = d['F_GAGAL'] & d['F_INVENTORI']
g = d.groupby(['DP Delivery', 'Kode Sprinter', 'Sprinter Delivery'])
r = pd.DataFrame({
    'Total Delivery': g.size(),
    'Total TTD': g['F_TTD'].sum(),
    'Total Scan Bermasalah': g['F_GAGAL_BM'].sum(),
    'Total AWB Tidak Scan Bermasalah': g['F_GAGAL_NO_BM'].sum(),
    'Total Scan Inventori': g['F_GAGAL_INV'].sum(),
    'Total AWB Tidak Scan Inventori': g['F_GAGAL_NO_INV'].sum(),
}).reset_index()
r['Keterangan'] = ''
r = r.sort_values(['DP Delivery', 'Total Delivery'], ascending=[True, False])
r.insert(0, 'No', range(1, len(r) + 1))

cols_num = ['Total Delivery', 'Total TTD', 'Total Scan Bermasalah', 'Total AWB Tidak Scan Bermasalah',
            'Total Scan Inventori', 'Total AWB Tidak Scan Inventori']
tot = {c: r[c].sum() for c in cols_num}
tot.update({'No': '', 'DP Delivery': 'TOTAL', 'Kode Sprinter': '', 'Sprinter Delivery': '', 'Keterangan': ''})
r = pd.concat([r, pd.DataFrame([tot])], ignore_index=True)
r = r[['No', 'DP Delivery', 'Kode Sprinter', 'Sprinter Delivery'] + cols_num + ['Keterangan']]

detail_cols = ['No. Waybill', 'DP Delivery', 'Kode Sprinter', 'Sprinter Delivery', 'Waktu Delivery',
               'Waktu Paket Bermasalah', 'Alasan Paket Bermasalah', 'Waktu Scan Inventory', 'Waktu TTD',
               'TTD Retur', 'DP TTD', 'Sprinter TTD', 'Tujuan', 'COD']
det = d[d['F_GAGAL_NO_INV']][detail_cols].sort_values(['Sprinter Delivery', 'Waktu Delivery'])
det_nobm = d[d['F_GAGAL_NO_BM']][detail_cols].sort_values(['Sprinter Delivery', 'Waktu Delivery'])

defs = pd.DataFrame({'Kolom': ['Total Delivery', 'Total TTD', 'Wajib Scan', 'Total Scan Bermasalah',
                               'Total AWB Tidak Scan Bermasalah', 'Total Scan Inventori',
                               'Total AWB Tidak Scan Inventori', 'Sumber'],
    'Definisi': ['Jumlah AWB yang di-scan delivery oleh sprinter',
                 'AWB TTD oleh penerima (Waktu TTD ada & TTD Retur = 0). TTD Retur = 1 tidak dihitung (itu TTD retur di pengirim)',
                 'Total Delivery - Total TTD = AWB yang wajib scan bermasalah dan wajib scan inventori',
                 'AWB wajib scan yang punya Waktu Paket Bermasalah',
                 'AWB wajib scan yang TIDAK punya Waktu Paket Bermasalah (Scan Bermasalah + Tidak Scan Bermasalah = Delivery - TTD)',
                 'AWB wajib scan yang punya Waktu Scan Inventory',
                 'AWB wajib scan yang TIDAK punya Waktu Scan Inventory (Scan Inventori + Tidak Scan Inventori = Delivery - TTD)',
                 '%s | Tanggal delivery: %s' % (os.path.basename(src), ', '.join(sorted({str(t) for t in tgl})))]})

os.makedirs(os.path.join(BASE, 'output'), exist_ok=True)
out = os.path.join(BASE, 'output', 'report_sprinter_%s.xlsx' % min(tgl).strftime('%Y%m%d'))
with pd.ExcelWriter(out, engine='openpyxl') as w:
    r.to_excel(w, sheet_name='Report Sprinter', index=False)
    det_nobm.to_excel(w, sheet_name='AWB Tidak Scan Bermasalah', index=False)
    det.to_excel(w, sheet_name='AWB Tidak Scan Inventori', index=False)
    defs.to_excel(w, sheet_name='Definisi', index=False)

wb = load_workbook(out)
thin = Side(style='thin', color='BFBFBF')
for ws in wb.worksheets:
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='C00000')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows():
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if hasattr(c.value, 'year'): c.number_format = 'yyyy-mm-dd hh:mm:ss'
    for i, col in enumerate(ws.columns, 1):
        ln = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(ln + 2, 10), 70)
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 32
ws = wb['Report Sprinter']
for row in ws.iter_rows(min_row=2):
    for k in (7, 9):
        if row[k].value and row[k].value > 0 and row[1].value != 'TOTAL':
            row[k].fill = PatternFill('solid', fgColor='FFC7CE')
last = ws.max_row
for c in ws[last]:
    c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='FFF2CC')
ws.column_dimensions['K'].width = 40
for c in ws['K'][1:]:
    c.alignment = Alignment(wrap_text=True, vertical='top')
wb['Definisi'].column_dimensions['B'].width = 110
wb.save(out)
print('OUTPUT:', out)
print(r.drop(columns=['Keterangan']).to_string(index=False))
